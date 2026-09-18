"""
Unit tests for Excel and CSV reporting module.
"""

from pathlib import Path
import openpyxl
import pytest
from database import init_db
from modules.projects import create_project
from modules.requirements import create_requirement
from modules.tasks import create_task
from modules.reports import generate_excel_report, generate_excel_bytes, export_table_csv


@pytest.fixture
def test_db(tmp_path):
    db_file = tmp_path / "reports_test.db"
    init_db(db_file)
    return db_file


def test_excel_report_generation_with_data(test_db, tmp_path):
    p_id = create_project({
        "project_code": "PRJ-REP-01",
        "project_name": "Report Validation Project",
        "project_type": "Equipment Inspection",
        "start_date": "2026-01-01",
        "deadline": "2026-06-01",
    }, db_path=test_db)

    create_requirement({
        "project_id": p_id,
        "requirement_code": "REQ-REP-101",
        "description": "Report test requirement",
    }, db_path=test_db)

    create_task({
        "project_id": p_id,
        "task_code": "TSK-REP-501",
        "task_description": "Report test task",
    }, db_path=test_db)

    out_file = tmp_path / "report_with_data.xlsx"
    generated_path = generate_excel_report(output_path=out_file, db_path=test_db)
    assert generated_path.exists()

    # Load workbook and verify sheets
    wb = openpyxl.load_workbook(generated_path)
    expected_sheets = [
        "Project Summary",
        "Requirements",
        "Tasks",
        "Deliverables",
        "Quality Checks",
        "Issues",
        "Activity Log",
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames

    # Check Project Summary has data
    ws = wb["Project Summary"]
    assert ws.max_row >= 2
    assert ws.cell(row=2, column=2).value == "PRJ-REP-01"


def test_excel_report_generation_empty_database(test_db, tmp_path):
    out_file = tmp_path / "report_empty.xlsx"
    generated_path = generate_excel_report(output_path=out_file, db_path=test_db)
    assert generated_path.exists()

    wb = openpyxl.load_workbook(generated_path)
    assert len(wb.sheetnames) == 7
    # Verify Project Summary sheet contains column headers and placeholder
    ws = wb["Project Summary"]
    assert ws.max_row >= 2
    assert "No records recorded" in str(ws.cell(row=2, column=1).value)


def test_generate_excel_bytes(test_db):
    b = generate_excel_bytes(db_path=test_db)
    assert isinstance(b, bytes)
    assert len(b) > 1000
    # First 4 bytes of a zip / xlsx file are PK\x03\x04
    assert b[:4] == b"PK\x03\x04"


def test_export_table_csv(test_db):
    p_id = create_project({
        "project_code": "PRJ-CSV",
        "project_name": "CSV Project",
        "project_type": "Equipment Inspection",
        "start_date": "2026-01-01",
        "deadline": "2026-06-01",
    }, db_path=test_db)

    csv_str = export_table_csv("projects", db_path=test_db)
    assert "PRJ-CSV" in csv_str
    assert "CSV Project" in csv_str

    # Non-existent table raises ValueError
    with pytest.raises(ValueError, match="Invalid table name"):
        export_table_csv("non_existent_table", db_path=test_db)
